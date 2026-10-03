"""Imports a GSC Job-to-Cash workbook (.xlsx, or its single-file CSV form; sheets Jobs, Agent_Tasks, Agents,
Rates_and_Details) into FlowGuard.

Safe to run again with a later workbook:
  * clients and agents are matched by name / Agent ID; master-data changes are applied and audited
  * jobs are matched by Agent Job Ref; new ones are added, changed ones are updated unless already invoiced
    (an invoiced job that changed at source is reported: it needs a credit/debit note, never a silent edit)
  * the first import creates the rate cards from the workbook; later imports never change prices. Any difference
    between the workbook's rates and the rates in force is reported for a governed configuration proposal.

Mapping (stored in the `workbook_categories` setting so it stays stable between imports):
  one FlowGuard job per agent task; the client's fixed job price sits on the completed task of the category's
  lead role (the role present on most jobs of that category); other roles on the job carry a £0 client side.
  Failed Visit / No Show tasks get their own job type with no rate, so they always go to a person.
"""
import io
import re
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal as D, ROUND_HALF_UP

from sqlalchemy import select

from .models import (Client, Agent, SelfBillingAgreement, Contract, JobType, RateCard, RateCardVersion, PricingRule,
                     VATConfiguration, Setting, Job, JobDecision, ImportBatch, Invoice)
from .services import processing
from .services import snapshot as snapmod
from .services.audit import log, notify

SOURCE = "GSC_WORKBOOK"
PENNY = D("0.01")
MAX_UNZIPPED = 100 * 1024 * 1024
SHEETS = ("Jobs", "Agent_Tasks", "Agents", "Rates_and_Details")


class WorkbookError(ValueError):
    pass


def money(v) -> D:
    return D(str(v).replace("£", "").replace(",", "").strip()).quantize(PENNY, rounding=ROUND_HALF_UP)


def dmy(s) -> date | None:
    if s in (None, ""):
        return None
    if isinstance(s, datetime):
        return s.date()
    if isinstance(s, date):
        return s
    s = str(s).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise WorkbookError(f"Unrecognised date '{s}'")


# ---------------------------------------------------------------------------------------------- reading
@dataclass
class Pack:
    jobs: dict
    tasks: list
    agents: dict
    clients: dict          # name -> {address, terms, ref_hint}
    client_prices: dict    # (client, category) -> D
    std_rates: dict        # role -> (rate, extra_per_30 | None)
    special_rates: dict    # (agent_id, role) -> (rate, extra_per_30 | None, note)
    company: dict          # name, address, vat_number
    numbering: dict        # INV/SB -> next number
    vat_rate: D | None


CSV_MARK = "#SHEET"


def _sheets_from_xlsx(source) -> dict:
    import openpyxl
    if isinstance(source, (bytes, bytearray)):
        try:
            zf = zipfile.ZipFile(io.BytesIO(source))
            if sum(i.file_size for i in zf.infolist()) > MAX_UNZIPPED:
                raise WorkbookError("Workbook expands to more than 100 MB")
        except zipfile.BadZipFile:
            raise WorkbookError("Not a valid .xlsx file")
        source = io.BytesIO(source)
    try:
        wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
        out = {name: [tuple(r) for r in wb[name].iter_rows(values_only=True)] for name in wb.sheetnames}
        wb.close()
        return out
    except WorkbookError:
        raise
    except Exception as e:  # openpyxl raises many types for malformed files
        raise WorkbookError(f"Could not read the workbook: {type(e).__name__}")


def _sheets_from_csv(data: bytes) -> dict:
    """The CSV form of the workbook: every sheet in one file, each starting with a '#SHEET,<name>' row
    followed by that sheet's own rows (see tools/workbook_to_csv.py). Empty cells become None."""
    import csv
    if b"\x00" in data:
        raise WorkbookError("The CSV file contains binary data")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("cp1252")  # saved from Excel on Windows
        except UnicodeDecodeError:
            raise WorkbookError("The CSV file is not UTF-8 text")
    out, cur = {}, None
    for row in csv.reader(io.StringIO(text)):
        if row and row[0].strip() == CSV_MARK:
            cur = row[1].strip() if len(row) > 1 else ""
            out[cur] = []
            continue
        if cur is None:
            if any(c.strip() for c in row):
                raise WorkbookError(f"This CSV is not a GSC data pack: it must start with a '{CSV_MARK},<sheet name>' row")
            continue
        out[cur].append(tuple(c.strip() if c.strip() != "" else None for c in row))
    return out


def read_pack(source) -> Pack:
    """source: a path, or the bytes of an .xlsx workbook or of its CSV form. Untrusted input: checked before use."""
    if isinstance(source, (bytes, bytearray)):
        sheets = _sheets_from_xlsx(source) if bytes(source[:2]) == b"PK" else _sheets_from_csv(bytes(source))
    elif str(source).lower().endswith(".csv"):
        with open(source, "rb") as f:
            sheets = _sheets_from_csv(f.read())
    else:
        sheets = _sheets_from_xlsx(source)
    for name, rows in sheets.items():  # pad ragged rows (CSV drops trailing empty cells) to the sheet's width
        width = max((len(r) for r in rows), default=0)
        sheets[name] = [tuple(r) + (None,) * (width - len(r)) for r in rows]
    missing = [s for s in SHEETS if s not in sheets]
    if missing:
        raise WorkbookError(f"Missing sheet(s): {', '.join(missing)}")

    def sheet(name, need, ints=()):
        rows = sheets[name]
        if not rows:
            raise WorkbookError(f"Sheet {name} is empty")
        head = [str(h).strip() if h is not None else "" for h in rows[0]]
        gone = [c for c in need if c not in head]
        if gone:
            raise WorkbookError(f"Sheet {name} is missing column(s): {', '.join(gone)}")
        out = [dict(zip(head, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
        for rec in out:  # whole-number columns arrive as text from CSV
            for c in ints:
                v = rec.get(c)
                if v not in (None, "") and not isinstance(v, int):
                    try:
                        rec[c] = int(float(str(v)))
                    except ValueError:
                        raise WorkbookError(f"Sheet {name}: '{v}' in column '{c}' is not a whole number")
        return out

    jobs = {j["GSC Job Ref"]: j for j in sheet("Jobs", ["GSC Job Ref", "Client Ref", "Client", "Job Category", "Postcode"])}
    tasks = sheet("Agent_Tasks", ["Agent Job Ref", "GSC Job Ref", "Client Ref", "Visit No", "Role", "Agent ID", "Visit Date",
                                  "Status", "Time On Site (mins)", "Verification", "Verification Date"],
                  ints=("Visit No", "Time On Site (mins)"))
    agents = {a["Agent ID"]: a for a in sheet("Agents", ["Agent ID", "Name", "Individual or Company", "Paid Via (Agent ID)",
                                                         "VAT Registered", "VAT Registered From", "VAT Number",
                                                         "Self-Billing Agreement", "Sends Own Invoices"])}
    section, header = "", None
    clients, prices, std, special, numbering, company = {}, {}, {}, {}, {}, {}
    vat_rate = None
    for r in sheets["Rates_and_Details"]:
        vals = [v for v in r if v not in (None, "")]
        if not vals:
            continue
        first = str(r[0] or "").strip()
        if len(vals) == 1 and r[0] not in (None, ""):
            section, header = first.lower(), None
            continue
        if header is None:
            header = r
            continue
        if section.startswith("gsc"):
            detail = str(r[1] or "")
            key = first.lower()
            if key == "company":
                company["name"] = re.sub(r"\s*\(.*\)\s*$", "", detail).strip()
            elif key == "address":
                company["address"] = detail
            elif key == "vat number":
                company["vat_number"] = detail
            elif key == "vat rate":
                m = re.search(r"(\d+(?:\.\d+)?)\s*%", detail)
                vat_rate = (D(m.group(1)) / 100) if m else None
                if vat_rate is not None:  # "0.20", "0.175": at least two decimal places, no trailing noise
                    vat_rate = vat_rate.normalize()
                    vat_rate = vat_rate.quantize(D("0.01")) if vat_rate.as_tuple().exponent > -2 else vat_rate
            elif key.startswith("next client invoice number") or key.startswith("next agent invoice number"):
                m = re.search(r"([A-Z]+)-(\d+)", detail)
                if m:
                    numbering[m.group(1)] = int(m.group(2))
        elif section.startswith("clients"):
            hint = str(r[3] or "")
            m = re.match(r"\s*([A-Z]{2,6})-", hint)
            clients[first] = {"address": r[1] or "", "terms": int(float(str(r[2]))) if r[2] not in (None, "") else 30,
                              "ref_prefix": m.group(1) if m else None}
        elif section.startswith("client prices"):
            prices[(first, str(r[1]).strip())] = money(r[2])
        elif section.startswith("agent standard"):
            std[first] = (money(r[1]), money(r[2]) if r[2] not in (None, "") else None)
        elif section.startswith("agent special"):
            special[(first, str(r[2]).strip())] = (money(r[3]), money(r[4]) if r[4] not in (None, "") else None, r[5] or "")
    if not prices or not std:
        raise WorkbookError("Rates_and_Details has no client prices or agent standard rates")
    return Pack(jobs, tasks, agents, clients, prices, std, special, company, numbering, vat_rate)


# ---------------------------------------------------------------------------------------------- mapping
def _abbrev(text: str, words_only=False) -> str:
    words = re.findall(r"[A-Za-z0-9]+", text)
    if not words:
        return "X"
    if not words_only and len(words) == 1 or (not words_only and re.search(r"\d", words[0])):
        return words[0].upper()[:6]
    return "".join(w[0] for w in words).upper()[:6]


def _unique(code: str, taken: set) -> str:
    out, n = code, 2
    while out in taken:
        out, n = f"{code[:5]}{n}", n + 1
    taken.add(out)
    return out


def category_map(db, pack: Pack) -> dict:
    """{category: {"code", "lead", "roles": {role: code}}}, merged with what earlier imports stored."""
    st = db.get(Setting, "workbook_categories")
    cfg = {k: {"code": v["code"], "lead": v["lead"], "roles": dict(v["roles"])} for k, v in ((st.value or {}) if st else {}).items()}
    role_counts: dict[str, Counter] = defaultdict(Counter)
    seen_jobs: dict[str, set] = defaultdict(set)
    for t in pack.tasks:
        cat = (pack.jobs.get(t["GSC Job Ref"]) or {}).get("Job Category")
        if cat and t["GSC Job Ref"] + "|" + t["Role"] not in seen_jobs[cat]:
            seen_jobs[cat].add(t["GSC Job Ref"] + "|" + t["Role"])
            role_counts[cat][t["Role"]] += 1
    cats = sorted({c for (_, c) in pack.client_prices} | set(role_counts))
    taken_cat = {v["code"] for v in cfg.values()}
    role_codes = {r: code for v in cfg.values() for r, code in v["roles"].items()}
    taken_role = set(role_codes.values())
    for cat in cats:
        if cat not in cfg:
            counts = role_counts.get(cat, Counter())
            lead = sorted(counts.items(), key=lambda x: (-x[1], x[0]))[0][0] if counts else None
            cfg[cat] = {"code": _unique(_abbrev(cat), taken_cat), "lead": lead, "roles": {}}
        for role in role_counts.get(cat, {}):
            if role not in cfg[cat]["roles"]:
                if role not in role_codes:
                    role_codes[role] = _unique(_abbrev(role, words_only=True), taken_role)
                cfg[cat]["roles"][role] = role_codes[role]
    if st:
        st.value = cfg
    else:
        db.add(Setting(key="workbook_categories", value=cfg))
    return cfg


def job_type_for(cfg: dict, category: str, role: str, status: str) -> str | None:
    c = cfg.get(category)
    if not c or role not in c["roles"]:
        return None
    code = f"{c['code']}.{c['roles'][role]}"
    return code + {"Failed Visit": ".FAIL", "No Show": ".NOSHOW"}.get(status, "")


def expected_rules(pack: Pack, cfg: dict, client_name: str) -> dict:
    """What the workbook says the rate card should contain: {(kind, job_type, agent_code): (client, agent, extra)}."""
    out = {}
    for (cname, category), price in pack.client_prices.items():
        if cname != client_name or category not in cfg:
            continue
        c = cfg[category]
        for role, rcode in c["roles"].items():
            if role not in pack.std_rates:
                continue
            jt = f"{c['code']}.{rcode}"
            rate, extra = pack.std_rates[role]
            out[("BASE", jt, None)] = (price if role == c["lead"] else D("0.00"), rate, extra)
            for (agent_code, srole), (srate, sextra, _) in pack.special_rates.items():
                if srole == role:
                    out[("AGENT_RATE", jt, agent_code)] = (D("0.00"), srate, sextra if sextra is not None else extra)
    return out


# ---------------------------------------------------------------------------------------------- import
def _set(db, key, value, actor, only_if_missing=False):
    st = db.get(Setting, key)
    if st and (only_if_missing or st.value == value):
        return False
    old = st.value if st else None
    if st:
        st.value = value
    else:
        db.add(Setting(key=key, value=value))
    log(db, actor, "SETTING_CHANGED", "setting", key, old=str(old)[:120] if old else "", new=str(value)[:120],
        reason="From imported workbook")
    return True


def import_workbook(db, actor: str, pack: Pack, filename: str = "workbook.xlsx", reprocess: bool = True) -> dict:
    s = {"filename": filename, "notes": [], "rate_differences": [], "changed_after_invoicing": [],
         "clients_added": 0, "agents_added": 0, "agents_updated": 0, "tasks_added": 0, "tasks_updated": 0,
         "tasks_unchanged": 0, "tasks_skipped": 0}
    today = snapmod.today()
    cfg = category_map(db, pack)

    # ---- company, VAT, numbering, self-billing policy
    if pack.company:
        _set(db, "company", pack.company, actor)
    vats = list(db.scalars(select(VATConfiguration).where(VATConfiguration.status == "ACTIVE")))
    if not vats and pack.vat_rate is not None:
        db.add(VATConfiguration(version=1, standard_rate=str(pack.vat_rate), valid_from=date(2000, 1, 1), status="ACTIVE",
                                note="Standard rate from the imported workbook."))
    elif vats and pack.vat_rate is not None and all(D(v.standard_rate) != pack.vat_rate for v in vats):
        s["notes"].append(f"Workbook VAT rate {pack.vat_rate * 100:.0f}% differs from the VAT configuration in force; not changed.")
    has_invoices = db.scalar(select(Invoice.id).limit(1)) is not None
    if pack.numbering and not has_invoices:
        _set(db, "invoice_numbering", pack.numbering, actor)
    _set(db, "self_billing_policy", {"sba_required_for": "VAT_REGISTERED"}, actor, only_if_missing=True)
    _set(db, "four_eyes_policy", {"mode": "ALL"}, actor, only_if_missing=True)

    # ---- clients and contracts
    clients = {c.name: c for c in db.scalars(select(Client))}
    taken = {c.code for c in clients.values()}
    for name, info in pack.clients.items():
        if name not in clients:
            code = _unique((info["ref_prefix"] or _abbrev(name, words_only=True))[:8], taken)
            c = Client(code=code, name=name, sector="", po_required=False, sage_account_ref=code, is_demo=False)
            db.add(c)
            db.flush()
            db.add(Contract(client_id=c.id, code=f"{code}-MAIN", version=1, valid_from=date(2000, 1, 1), status="ACTIVE",
                            terms={"payment_days": info["terms"], "billing_address": info["address"]}))
            clients[name] = c
            s["clients_added"] += 1
            log(db, actor, "CLIENT_CREATED", "client", code, new=name, reason="From imported workbook")
        else:
            con = db.scalar(select(Contract).where(Contract.client_id == clients[name].id, Contract.status == "ACTIVE"))
            if con and (con.terms or {}).get("payment_days") != info["terms"]:
                con.terms = dict(con.terms or {}, payment_days=info["terms"], billing_address=info["address"])
                log(db, actor, "CONTRACT_TERMS_UPDATED", "contract", con.code, new=f"{info['terms']} days")

    # ---- agents (master data from the source system)
    agents = {a.code: a for a in db.scalars(select(Agent))}
    for aid, a in pack.agents.items():
        vals = {"name": a["Name"], "supplier_type": "LIMITED_COMPANY" if a["Individual or Company"] == "Company" else "SOLE_TRADER",
                "vat_status": {"Yes": "REGISTERED", "No": "NOT_REGISTERED"}.get(a["VAT Registered"], "UNKNOWN"),
                "vat_number": a["VAT Number"] or None, "self_billing": a["Sends Own Invoices"] != "Yes",
                "region": a.get("Address") or "", "vat_registered_from": dmy(a["VAT Registered From"])}
        if aid not in agents:
            ag = Agent(code=aid, sage_supplier_ref=aid.replace("-", ""), is_demo=False, **vals)
            db.add(ag)
            agents[aid] = ag
            s["agents_added"] += 1
        else:
            ag = agents[aid]
            diff = {k: v for k, v in vals.items() if getattr(ag, k) != v}
            if diff:
                for k, v in diff.items():
                    setattr(ag, k, v)
                s["agents_updated"] += 1
                log(db, actor, "AGENT_UPDATED", "agent", aid, reason="From imported workbook: " + ", ".join(diff),
                    details={k: str(v) for k, v in diff.items()})
    db.flush()
    for aid, a in pack.agents.items():
        ag = agents[aid]
        via = agents.get(a["Paid Via (Agent ID)"]).id if a["Paid Via (Agent ID)"] in agents else None
        if a["Paid Via (Agent ID)"] and not via:
            s["notes"].append(f"{aid} is paid via unknown agent {a['Paid Via (Agent ID)']}.")
        if ag.paid_via_agent_id != via:
            ag.paid_via_agent_id = via
        sba = db.scalar(select(SelfBillingAgreement).where(SelfBillingAgreement.agent_id == ag.id,
                                                           SelfBillingAgreement.status == "ACTIVE"))
        if a["Self-Billing Agreement"] == "Yes" and not sba:
            db.add(SelfBillingAgreement(agent_id=ag.id, reference=f"SBA-{aid}", valid_from=date(2000, 1, 1), status="ACTIVE"))
            log(db, actor, "SBA_RECORDED", "agent", aid, new="ACTIVE", reason="Workbook: self-billing agreement Yes (start date not given)")
        elif a["Self-Billing Agreement"] == "No" and sba:
            sba.status, sba.valid_to = "ENDED", today
            log(db, actor, "SBA_ENDED", "agent", aid, old="ACTIVE", new="ENDED", reason="Workbook: self-billing agreement No")

    # ---- job types and rate cards
    for cat, c in cfg.items():
        for role, rcode in c["roles"].items():
            for suffix, label in (("", ""), (".FAIL", " (failed visit)"), (".NOSHOW", " (no show)")):
                code = f"{c['code']}.{rcode}{suffix}"
                if not db.get(JobType, code):
                    db.add(JobType(code=code, name=f"{cat} - {role}{label}"[:120]))
    first_task = min((dmy(t["Visit Date"]) for t in pack.tasks if t.get("Visit Date")), default=today)
    start = first_task.replace(day=1)
    n_rule = [max([int(c[1:]) for (c,) in db.execute(select(PricingRule.rule_code)) if c[1:].isdigit()] + [0])]
    for name, c in clients.items():
        exp = expected_rules(pack, cfg, name)
        if not exp:
            continue
        card = db.scalar(select(RateCard).where(RateCard.client_id == c.id))
        if not card:
            card = RateCard(client_id=c.id, contract_code=f"{c.code}-MAIN", name=f"{c.code} rates")
            db.add(card)
            db.flush()
            v = RateCardVersion(rate_card_id=card.id, version=1, valid_from=start, status="ACTIVE", precedence=[],
                                note=f"Created from {filename}", created_by=actor, approved_by=f"{actor} (initial import)")
            db.add(v)
            db.flush()
            for (kind, jt, agent_code), (ca, aa, extra) in sorted(exp.items(), key=lambda x: (x[0][1], x[0][0], x[0][2] or "")):
                n_rule[0] += 1
                db.add(PricingRule(rcv_id=v.id, rule_code=f"R{n_rule[0]}", kind=kind, job_type=jt, client_amount=ca,
                                   agent_amount=aa, agent_code=agent_code, included_minutes=60 if extra is not None else None,
                                   agent_extra_per_30=extra,
                                   clause_ref="Workbook client price + agent rate" if kind == "BASE" else "Workbook special agent rate"))
            log(db, actor, "RATE_CARD_VERSION_CREATED", "rate_card_version", v.id, new=f"{card.name} v1",
                reason=f"Initial rates from {filename}", config_version="v1")
            continue
        # existing configuration: compare, never change
        latest = max((v for v in card.versions if v.status == "ACTIVE"), key=lambda v: v.valid_from, default=None)
        have = {}
        for r in (latest.rules if latest else []):
            if r.kind in ("BASE", "AGENT_RATE"):
                have[(r.kind, r.job_type, r.agent_code)] = (r.client_amount, r.agent_amount, r.agent_extra_per_30)
        for k, want in sorted(exp.items(), key=lambda x: (x[0][1], x[0][0], x[0][2] or "")):
            got = have.get(k)
            if got != want:
                s["rate_differences"].append({"client": c.code, "kind": k[0], "job_type": k[1], "agent": k[2],
                                              "in_force": [str(x) if x is not None else None for x in got] if got else None,
                                              "workbook": [str(x) if x is not None else None for x in want]})

    # ---- tasks -> jobs
    existing = {j.external_job_id: j for j in db.scalars(select(Job).where(Job.source_system == SOURCE))}
    frozen = {jid for (jid,) in db.execute(select(JobDecision.job_id).where(JobDecision.frozen == True))}  # noqa: E712
    for t in pack.tasks:
        gj = pack.jobs.get(t["GSC Job Ref"])
        if not gj:
            s["tasks_skipped"] += 1
            s["notes"].append(f"{t['Agent Job Ref']}: GSC job {t['GSC Job Ref']} is not on the Jobs sheet; skipped.")
            continue
        c, a = clients.get(gj["Client"]), agents.get(t["Agent ID"])
        ver = {"Verified": "VERIFIED", "Rejected": "REJECTED"}.get(t["Verification"], "UNVERIFIED")
        vdate, visit = dmy(t["Verification Date"]), dmy(t["Visit Date"])
        pa = pack.agents.get(t["Agent ID"], {})
        tags = [x for x in (t["Status"] if t["Status"] != "Completed" else "",
                            "revisit" if (t["Visit No"] or 1) > 1 else "",
                            "sub-agent" if pa.get("Paid Via (Agent ID)") else "") if x]
        vals = {"client_id": c.id if c else None, "agent_id": a.id if a else None,
                "job_type": job_type_for(cfg, gj["Job Category"], t["Role"], t["Status"]), "job_date": visit,
                "verification_status": ver,
                "verification_timestamp": datetime.combine(vdate, time(17, 0)) if vdate and ver != "UNVERIFIED" else None,
                "verified_by": "GSC admin (workbook)" if ver != "UNVERIFIED" else None,
                "weekend": bool(visit and visit.weekday() >= 5), "revisit": (t["Visit No"] or 1) > 1,
                "postcode": gj.get("Postcode"), "minutes_on_site": t["Time On Site (mins)"],
                "notes": f"Client ref {t['Client Ref']} · {gj.get('Customer Name') or '-'} · {t['Role']} · visit "
                         f"{t['Visit No']} · {t['Status']}" + (f" · team: {gj['Team Set-up']}" if gj.get("Team Set-up") else ""),
                "source_record_id": t["GSC Job Ref"], "demo_scenario": ", ".join(tags)[:60]}
        j = existing.get(t["Agent Job Ref"])
        if j is None:
            j = Job(external_job_id=t["Agent Job Ref"], source_system=SOURCE, emergency=False, **vals)
            db.add(j)
            existing[j.external_job_id] = j
            s["tasks_added"] += 1
            continue
        diff = {k: v for k, v in vals.items() if getattr(j, k) != v and k not in ("notes", "demo_scenario")}
        if not diff:
            s["tasks_unchanged"] += 1
            continue
        if j.id in frozen:
            s["changed_after_invoicing"].append({"task": j.external_job_id, "fields": sorted(diff)})
            continue
        for k, v in vals.items():
            setattr(j, k, v)
        j.updated_at = datetime.utcnow().replace(microsecond=0)
        s["tasks_updated"] += 1
        log(db, actor, "JOB_UPDATED_FROM_SOURCE", "job", j.id, reason=f"{filename}: " + ", ".join(sorted(diff)),
            details={k: str(v) for k, v in diff.items()})
    db.flush()

    future = sum(1 for t in pack.tasks if t.get("Visit Date") and dmy(t["Visit Date"]) > today)
    if future:
        s["notes"].append(f"{future} tasks are dated after the business date ({today:%d/%m/%Y}). They will show as "
                          "'invalid date' until the business date is moved on (Settings → Business date).")
    if s["changed_after_invoicing"]:
        s["notes"].append(f"{len(s['changed_after_invoicing'])} already-invoiced tasks changed at source. Nothing was "
                          "changed in FlowGuard; these need a credit/debit note.")
    if s["rate_differences"]:
        s["notes"].append(f"{len(s['rate_differences'])} rate differences between the workbook and the rates in force. "
                          "Prices were not changed: raise a configuration proposal if the workbook is right.")
    batch = ImportBatch(filename=filename[:120], uploaded_by=actor, status="IMPORTED_LIVE", headers=list(SHEETS), mapping={},
                        summary={k: v for k, v in s.items() if k not in ("rate_differences", "changed_after_invoicing")}
                        | {"kind": "WORKBOOK", "rate_differences": s["rate_differences"][:200],
                           "changed_after_invoicing": s["changed_after_invoicing"][:200]})
    db.add(batch)
    log(db, actor, "WORKBOOK_IMPORTED", "import_batch", filename[:48],
        new=f"{s['tasks_added']} added, {s['tasks_updated']} updated",
        details={k: v for k, v in s.items() if isinstance(v, int)})
    db.commit()
    s["batch_id"] = batch.id
    if reprocess:
        res = processing.process(db, actor=actor)
        s["counts"] = res["counts"]
        batch.summary = dict(batch.summary, counts=res["counts"])
        notify(db, "CSV_IMPORT", f"Workbook imported: {filename}",
               f"{s['tasks_added']} new tasks, {s['tasks_updated']} updated. Ready {res['counts']['READY']}, "
               f"review {res['counts']['NEEDS_REVIEW']}, blocked {res['counts']['BLOCKED']}.")
        db.commit()
    return s
